import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
from pydantic import BaseModel

class ParsedElement(BaseModel):
    page_number: int
    element_type: str  # "heading", "paragraph", "code_block", "table", "procedure"
    content: str
    section: str = "General"
    parent_section: str = ""
    metadata: Dict[str, Any] = {}

class ParsedDocument(BaseModel):
    filename: str
    file_type: str
    file_size: int
    page_count: int
    elements: List[ParsedElement]
    raw_text: str = ""

class DocumentParser:
    """Multi-format parser for technical documentation."""

    SUPPORTED_EXTENSIONS = {
        ".pdf", ".docx", ".txt", ".md",
        ".js", ".jsx", ".ts", ".tsx",
        ".py", ".sql", ".html", ".css",
        ".json", ".yaml", ".yml", ".sh",
        ".java", ".cs", ".cpp", ".c"
    }

    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def parse_file(cls, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        file_size = path.stat().st_size

        if ext == ".pdf":
            return cls._parse_pdf(path, file_size)
        elif ext == ".docx":
            return cls._parse_docx(path, file_size)
        elif ext in [".md", ".markdown"]:
            return cls._parse_markdown(path, file_size)
        elif ext in [".js", ".jsx", ".ts", ".tsx", ".py", ".sql", ".html", ".css", ".json", ".yaml", ".yml", ".sh", ".java", ".cs", ".cpp", ".c"]:
            return cls._parse_code_file(path, file_size, ext)
        else:
            return cls._parse_plain_text(path, file_size)

    @classmethod
    def _parse_pdf(cls, path: Path, file_size: int) -> ParsedDocument:
        import pypdf
        elements = []
        full_text = []

        reader = pypdf.PdfReader(str(path))
        page_count = len(reader.pages)

        current_section = "Introduction"
        parent_section = ""

        for idx, page in enumerate(reader.pages):
            page_num = idx + 1
            text = page.extract_text() or ""
            if not text.strip():
                continue

            full_text.append(f"--- Page {page_num} ---\n" + text)
            lines = text.splitlines()

            buffer = []
            for line in lines:
                stripped = line.strip()
                if not stripped:
                    if buffer:
                        elements.append(ParsedElement(
                            page_number=page_num,
                            element_type="paragraph",
                            content="\n".join(buffer),
                            section=current_section,
                            parent_section=parent_section
                        ))
                        buffer = []
                    continue

                # Heuristic for headings in PDF: short line, capitalized, or numbered
                if (len(stripped) < 70 and 
                    (stripped.isupper() or 
                     any(stripped.startswith(prefix) for prefix in ["Section ", "Chapter ", "#"]) or
                     (stripped[0].isdigit() and "." in stripped[:4]))):
                    
                    if buffer:
                        elements.append(ParsedElement(
                            page_number=page_num,
                            element_type="paragraph",
                            content="\n".join(buffer),
                            section=current_section,
                            parent_section=parent_section
                        ))
                        buffer = []
                    
                    parent_section = current_section
                    current_section = stripped.lstrip("#").strip()
                    elements.append(ParsedElement(
                        page_number=page_num,
                        element_type="heading",
                        content=current_section,
                        section=current_section,
                        parent_section=parent_section
                    ))
                else:
                    buffer.append(stripped)

            if buffer:
                elements.append(ParsedElement(
                    page_number=page_num,
                    element_type="paragraph",
                    content="\n".join(buffer),
                    section=current_section,
                    parent_section=parent_section
                ))

        return ParsedDocument(
            filename=path.name,
            file_type="pdf",
            file_size=file_size,
            page_count=max(page_count, 1),
            elements=elements,
            raw_text="\n\n".join(full_text)
        )

    @classmethod
    def _parse_docx(cls, path: Path, file_size: int) -> ParsedDocument:
        import docx
        doc = docx.Document(str(path))
        elements = []
        full_text = []

        current_section = "Overview"
        parent_section = ""
        current_page = 1
        para_count = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            para_count += 1
            # Approximate page progression (approx 15-20 paragraphs per page)
            if para_count % 18 == 0:
                current_page += 1

            full_text.append(text)
            style_name = para.style.name.lower() if para.style else ""

            if "heading" in style_name or "title" in style_name:
                parent_section = current_section
                current_section = text
                elements.append(ParsedElement(
                    page_number=current_page,
                    element_type="heading",
                    content=text,
                    section=current_section,
                    parent_section=parent_section
                ))
            else:
                elements.append(ParsedElement(
                    page_number=current_page,
                    element_type="paragraph",
                    content=text,
                    section=current_section,
                    parent_section=parent_section
                ))

        # Also extract tables from docx
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            if table_rows:
                table_str = "\n".join(table_rows)
                elements.append(ParsedElement(
                    page_number=current_page,
                    element_type="table",
                    content=table_str,
                    section=current_section,
                    parent_section=parent_section
                ))

        return ParsedDocument(
            filename=path.name,
            file_type="docx",
            file_size=file_size,
            page_count=max(current_page, 1),
            elements=elements,
            raw_text="\n\n".join(full_text)
        )

    @classmethod
    def _parse_markdown(cls, path: Path, file_size: int) -> ParsedDocument:
        content = path.read_text(encoding="utf-8", errors="replace")
        lines = content.splitlines()

        elements = []
        current_section = "Overview"
        parent_section = ""
        current_page = 1

        buffer = []
        in_code_block = False
        code_lang = ""
        code_buffer = []

        for line_idx, line in enumerate(lines):
            # Check for page markers like <!-- page 10 --> or --- Page 10 ---
            page_match = re.search(r"(?:<!--\s*page\s*(\d+)\s*-->|---\s*Page\s*(\d+)\s*---)", line, re.IGNORECASE)
            if page_match:
                num_str = page_match.group(1) or page_match.group(2)
                current_page = int(num_str)
                continue
            elif "<!-- page" in line.lower() or ("page" in line.lower() and line.strip().startswith("---")):
                current_page += 1
                continue

            stripped = line.strip()

            if stripped.startswith("```"):
                if not in_code_block:
                    # Flush paragraph buffer
                    if buffer:
                        elements.append(ParsedElement(
                            page_number=current_page,
                            element_type="paragraph",
                            content="\n".join(buffer),
                            section=current_section,
                            parent_section=parent_section
                        ))
                        buffer = []
                    in_code_block = True
                    code_lang = stripped.lstrip("`").strip()
                    code_buffer = []
                else:
                    in_code_block = False
                    elements.append(ParsedElement(
                        page_number=current_page,
                        element_type="code_block",
                        content="\n".join(code_buffer),
                        section=current_section,
                        parent_section=parent_section,
                        metadata={"language": code_lang}
                    ))
                    code_buffer = []
                continue

            if in_code_block:
                code_buffer.append(line)
                continue

            if stripped.startswith("#"):
                if buffer:
                    elements.append(ParsedElement(
                        page_number=current_page,
                        element_type="paragraph",
                        content="\n".join(buffer),
                        section=current_section,
                        parent_section=parent_section
                    ))
                    buffer = []

                heading_text = stripped.lstrip("#").strip()
                parent_section = current_section
                current_section = heading_text
                elements.append(ParsedElement(
                    page_number=current_page,
                    element_type="heading",
                    content=heading_text,
                    section=current_section,
                    parent_section=parent_section
                ))
            elif stripped.startswith("|") and stripped.endswith("|"):
                buffer.append(line)
            elif not stripped:
                if buffer:
                    # check if buffer is table
                    is_table = all(b.strip().startswith("|") for b in buffer)
                    elements.append(ParsedElement(
                        page_number=current_page,
                        element_type="table" if is_table else "paragraph",
                        content="\n".join(buffer),
                        section=current_section,
                        parent_section=parent_section
                    ))
                    buffer = []
            else:
                buffer.append(line)

        if buffer:
            elements.append(ParsedElement(
                page_number=current_page,
                element_type="paragraph",
                content="\n".join(buffer),
                section=current_section,
                parent_section=parent_section
            ))

        return ParsedDocument(
            filename=path.name,
            file_type="markdown",
            file_size=file_size,
            page_count=max(current_page, 1),
            elements=elements,
            raw_text=content
        )

    @classmethod
    def _parse_code_file(cls, path: Path, file_size: int, ext: str) -> ParsedDocument:
        content = path.read_text(encoding="utf-8", errors="replace")
        lang_map = {
            ".js": "javascript", ".jsx": "javascript",
            ".ts": "typescript", ".tsx": "typescript",
            ".py": "python", ".sql": "sql",
            ".html": "html", ".css": "css",
            ".json": "json", ".yaml": "yaml", ".yml": "yaml",
            ".sh": "bash", ".java": "java", ".cs": "csharp",
            ".cpp": "cpp", ".c": "c"
        }
        lang = lang_map.get(ext, "code")

        elements = [
            ParsedElement(
                page_number=1,
                element_type="code_block",
                content=content,
                section=f"{path.stem} Source File",
                parent_section="",
                metadata={"language": lang}
            )
        ]

        return ParsedDocument(
            filename=path.name,
            file_type="code",
            file_size=file_size,
            page_count=1,
            elements=elements,
            raw_text=content
        )

    @classmethod
    def _parse_plain_text(cls, path: Path, file_size: int) -> ParsedDocument:
        content = path.read_text(encoding="utf-8", errors="replace")
        lines = content.splitlines()
        elements = []
        current_section = "General"
        buffer = []
        page_number = 1

        for line in lines:
            stripped = line.strip()
            if "page " in stripped.lower() and len(stripped) < 20:
                p_match = re.search(r"page\s*(\d+)", stripped, re.IGNORECASE)
                if p_match:
                    page_number = int(p_match.group(1))
                else:
                    page_number += 1
                continue
            if not stripped:
                if buffer:
                    elements.append(ParsedElement(
                        page_number=page_number,
                        element_type="paragraph",
                        content="\n".join(buffer),
                        section=current_section,
                        parent_section=""
                    ))
                    buffer = []
            else:
                buffer.append(line)

        if buffer:
            elements.append(ParsedElement(
                page_number=page_number,
                element_type="paragraph",
                content="\n".join(buffer),
                section=current_section,
                parent_section=""
            ))

        return ParsedDocument(
            filename=path.name,
            file_type="txt",
            file_size=file_size,
            page_count=max(page_number, 1),
            elements=elements,
            raw_text=content
        )
